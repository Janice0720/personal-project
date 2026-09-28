"""Insert a formatted cover page at the very start of the exported docx
(before pandoc's TOC block), matching 國立臺北商業大學資訊與決策科學研究所 format spec.

Run this AFTER the pandoc conversion step (see README.md), since it edits the
already-generated docx in place.
"""
import os
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUTPUT_FILENAME = "黃子甄_論文計畫書(第一至三章初稿).docx"
DOCX_PATH = os.path.normpath(os.path.join(os.path.dirname(__file__), "..", OUTPUT_FILENAME))

# ---- 封面資訊：如有異動請直接修改以下欄位後重新執行本腳本 ----
SCHOOL = "國立臺北商業大學"
INSTITUTE = "資訊與決策科學研究所"
DEGREE_LABEL = "碩士學位論文"
TITLE_ZH = "應用投資人分群與個人化推薦模型於數位券商基金投資服務之研究"
TITLE_EN = "A Study on Applying Investor Segmentation and Personalized Recommendation Models to Digital Brokerage Fund Investment Services"
STUDENT = "研究生：黃子甄　撰"
ADVISOR = "指導教授：王亦凡　博士"
GRAD_DATE = "中華民國一一六年六月"  # 115 學年度第二學期（116 年 6 月）預計畢業


def set_run_font(run, size, bold, east_asia='標楷體', ascii_font='Times New Roman'):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = ascii_font
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), east_asia)
    rFonts.set(qn('w:ascii'), ascii_font)
    rFonts.set(qn('w:hAnsi'), ascii_font)


def make_para(doc, text, size, bold, east_asia='標楷體', ascii_font='Times New Roman',
              space_before=0, space_after=0, line_spacing_pt=None, page_break=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    if line_spacing_pt:
        pf.line_spacing = Pt(line_spacing_pt)
    run = p.add_run(text if text else "")
    set_run_font(run, size, bold, east_asia, ascii_font)
    if page_break:
        run.add_break(WD_BREAK.PAGE)
    return p


def main():
    doc = Document(DOCX_PATH)
    body = doc.element.body
    first_child = body[0]  # pandoc puts the TOC content-control (or first para) here

    cover_paras = []
    cover_paras.append(make_para(doc, SCHOOL, 24, True, space_before=40))
    cover_paras.append(make_para(doc, INSTITUTE, 24, True))
    cover_paras.append(make_para(doc, DEGREE_LABEL, 24, True))
    for _ in range(5):
        cover_paras.append(make_para(doc, "", 18, False))
    cover_paras.append(make_para(doc, TITLE_ZH, 24, True, line_spacing_pt=32))
    cover_paras.append(make_para(doc, "", 12, False))
    cover_paras.append(make_para(doc, TITLE_EN, 20, False, line_spacing_pt=26))
    for _ in range(5):
        cover_paras.append(make_para(doc, "", 18, False))
    cover_paras.append(make_para(doc, STUDENT, 18, True))
    for _ in range(4):
        cover_paras.append(make_para(doc, "", 18, False))
    cover_paras.append(make_para(doc, ADVISOR, 18, True))
    for _ in range(4):
        cover_paras.append(make_para(doc, "", 18, False))
    cover_paras.append(make_para(doc, GRAD_DATE, 18, True))
    page_break_para = doc.add_paragraph()
    page_break_para.add_run().add_break(WD_BREAK.PAGE)
    cover_paras.append(page_break_para)

    # Each paragraph above was appended at the very end of the body; move them,
    # in the same order, to sit immediately before the original first element
    # (pandoc's TOC block), so the final order is: cover -> TOC -> chapter 1.
    insert_ref = first_child
    for p in cover_paras:
        el = p._p
        body.remove(el)
        insert_ref.addprevious(el)

    doc.save(DOCX_PATH)
    print("cover page inserted before TOC ->", DOCX_PATH)


if __name__ == "__main__":
    main()
